#!/usr/bin/env python3
"""CMMC Evidence Prefix Manager.

Copies evidence into a new folder and prefixes each filename using an imported
Document Request List (DRL): <ObjectiveID>_<ERL#>_<original filename>.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import os
import re
import shutil
import sys
import threading
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable


OBJECTIVE_ALIASES = {
    "objectiveid", "objective", "assessmentobjective", "controlobjective",
    "objectiveidentifier", "nistobjectiveid", "cmmcobjectiveid",
}
ERL_ALIASES = {
    "erl", "erlid", "erlnumber", "erlno", "documentrequestid",
    "evidencerequestid", "requestid", "requestnumber",
}
DOMAIN_ALIASES = {"domain", "family", "securitydomain", "controlfamily"}
CONTROL_ALIASES = {
    "control", "controlid", "requirement", "requirementid", "practice",
    "practiceid",
}
DESCRIPTION_ALIASES = {
    "description", "request", "documentrequest", "evidencerequest",
    "artifact", "requestedartifact", "title", "name", "securityrequirement",
}
OBJECTIVE_RE = re.compile(
    r"(?i)(?<![0-9A-Z])(?:[A-Z]{2}\.L2[-_ ]*)?3[._ -]\d{1,2}[._ -]\d{1,2}"
    r"(?:\s*[\[(._ -]?\s*[A-Z]\s*[\])]?)?(?![0-9A-Z])"
)


def header_key(value: object) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def clean_cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def safe_component(value: str) -> str:
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "-", value.strip())
    value = re.sub(r"\s+", " ", value).rstrip(". ")
    return value


def normalized_token(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", value.casefold())


def canonical_objective(value: str) -> str:
    """Normalize punctuation/case while preserving an optional CMMC family prefix."""
    raw = value.strip()
    match = re.search(
        r"(?i)(?:([A-Z]{2})\.L2[-_ ]*)?(3)[._ -](\d{1,2})[._ -](\d{1,2})"
        r"(?:\s*[\[(._ -]?\s*([A-Z])\s*[\])]?)?",
        raw,
    )
    if not match:
        return safe_component(raw)
    family, major, minor, req, letter = match.groups()
    core = f"{major}.{int(minor)}.{int(req)}"
    if letter:
        core += f"[{letter.lower()}]"
    return f"{family.upper()}.L2-{core}" if family else core


@dataclass(frozen=True)
class MappingRow:
    objective: str
    erl: str
    domain: str = ""
    control: str = ""
    description: str = ""
    row_number: int = 0

    @property
    def label(self) -> str:
        return f"{self.objective} | {self.erl}"


@dataclass
class FilePlan:
    source: Path
    relative_path: Path
    mapping: MappingRow | None
    status: str
    match_basis: str
    candidates: list[MappingRow] = field(default_factory=list)
    new_name: str = ""
    note: str = ""


def find_column(headers: list[str], aliases: set[str]) -> int | None:
    keys = [header_key(h) for h in headers]
    for i, key in enumerate(keys):
        if key in aliases:
            return i
    # Accept explanatory text appended to a recognized header, such as
    # "ERL # (Suggested identifier for the evidence)". Exact matches above
    # retain priority when a workbook contains similarly named columns.
    for alias in sorted(aliases, key=len, reverse=True):
        if len(alias) < 3:
            continue
        for i, key in enumerate(keys):
            if key.startswith(alias):
                return i
    return None


def _rows_from_delimited(path: Path) -> list[list[object]]:
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return [list(row) for row in csv.reader(handle, delimiter=delimiter)]


def _rows_from_xlsx(path: Path, sheet_name: str | None = None) -> list[list[object]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Excel input requires openpyxl. Run: pip install openpyxl") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[sheet_name] if sheet_name else workbook.active
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def load_mapping(path: Path, sheet_name: str | None = None) -> list[MappingRow]:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        rows = _rows_from_xlsx(path, sheet_name)
    elif suffix in {".csv", ".tsv"}:
        rows = _rows_from_delimited(path)
    else:
        raise ValueError("Mapping file must be CSV, TSV, XLSX, or XLSM.")
    rows = [row for row in rows if any(clean_cell(v) for v in row)]
    if not rows:
        raise ValueError("The mapping file is empty.")

    header_index = None
    objective_col = erl_col = None
    for i, row in enumerate(rows[:25]):
        headers = [clean_cell(v) for v in row]
        obj = find_column(headers, OBJECTIVE_ALIASES)
        erl = find_column(headers, ERL_ALIASES)
        if obj is not None and erl is not None:
            header_index, objective_col, erl_col = i, obj, erl
            break
    if header_index is None or objective_col is None or erl_col is None:
        raise ValueError(
            "Could not locate Objective ID and ERL columns in the first 25 rows. "
            "Use headers such as 'Objective ID' and 'ERL #'."
        )

    headers = [clean_cell(v) for v in rows[header_index]]
    domain_col = find_column(headers, DOMAIN_ALIASES)
    control_col = find_column(headers, CONTROL_ALIASES)
    description_col = find_column(headers, DESCRIPTION_ALIASES)

    def at(row: list[object], index: int | None) -> str:
        return clean_cell(row[index]) if index is not None and index < len(row) else ""

    mappings: list[MappingRow] = []
    seen: set[tuple[str, str]] = set()
    for excel_row, row in enumerate(rows[header_index + 1 :], header_index + 2):
        objective = canonical_objective(at(row, objective_col))
        erl = safe_component(at(row, erl_col))
        if not objective or not erl:
            continue
        key = (normalized_token(objective), normalized_token(erl))
        if key in seen:
            continue
        seen.add(key)
        mappings.append(
            MappingRow(
                objective=objective,
                erl=erl,
                domain=at(row, domain_col),
                control=at(row, control_col),
                description=at(row, description_col),
                row_number=excel_row,
            )
        )
    if not mappings:
        raise ValueError("No complete Objective ID and ERL mapping rows were found.")
    return mappings


def boundary_contains(text: str, token: str) -> bool:
    if not token:
        return False
    pattern = r"(?<![A-Za-z0-9])" + re.escape(token) + r"(?![A-Za-z0-9])"
    return re.search(pattern, text, flags=re.IGNORECASE) is not None


def fuzzy_contains(text: str, token: str) -> bool:
    """Match tokens despite common punctuation differences, with digit boundaries."""
    needle = normalized_token(token)
    if not needle:
        return False
    haystack = normalized_token(text)
    start = 0
    while True:
        index = haystack.find(needle, start)
        if index < 0:
            return False
        before = haystack[index - 1] if index else ""
        after_pos = index + len(needle)
        after = haystack[after_pos] if after_pos < len(haystack) else ""
        if not (needle[0].isdigit() and before.isdigit()) and not (
            needle[-1].isdigit() and after.isdigit()
        ):
            return True
        start = index + 1


def match_file(path: Path, relative_path: Path, mappings: list[MappingRow]) -> FilePlan:
    searchable = str(relative_path.with_suffix(""))
    objective_hits = [m for m in mappings if fuzzy_contains(searchable, m.objective)]
    erl_hits = [m for m in mappings if boundary_contains(searchable, m.erl) or fuzzy_contains(searchable, m.erl)]
    both = [m for m in objective_hits if m in erl_hits]

    if len(both) == 1:
        chosen, basis = both[0], "Objective ID + ERL"
        return build_plan(path, relative_path, chosen, "Ready", basis, both)
    if len(both) > 1:
        return build_plan(path, relative_path, None, "Ambiguous", "Multiple Objective ID + ERL matches", both)
    if len(objective_hits) == 1:
        chosen = objective_hits[0]
        return build_plan(path, relative_path, chosen, "Ready", "Objective ID", objective_hits)
    if len(erl_hits) == 1:
        chosen = erl_hits[0]
        return build_plan(path, relative_path, chosen, "Ready", "ERL", erl_hits)

    candidates = list(dict.fromkeys(objective_hits + erl_hits))
    if candidates:
        # Domain/control folder names can resolve a repeated ERL or objective.
        scored: list[tuple[int, MappingRow]] = []
        for item in candidates:
            score = sum(
                1 for hint in (item.domain, item.control)
                if hint and fuzzy_contains(searchable, hint)
            )
            scored.append((score, item))
        best_score = max(score for score, _ in scored)
        best = [item for score, item in scored if score == best_score]
        if best_score > 0 and len(best) == 1:
            return build_plan(path, relative_path, best[0], "Ready", "Folder/domain hint", candidates)
        reason = "Multiple possible mapping rows"
        return build_plan(path, relative_path, None, "Ambiguous", reason, candidates)
    return build_plan(path, relative_path, None, "Unmatched", "No Objective ID or ERL match", [])


def build_plan(
    source: Path,
    relative_path: Path,
    mapping: MappingRow | None,
    status: str,
    basis: str,
    candidates: list[MappingRow],
) -> FilePlan:
    new_name = prefixed_name(source.name, mapping) if mapping else ""
    if mapping and source.name.casefold() == new_name.casefold():
        status = "Already aligned"
    return FilePlan(source, relative_path, mapping, status, basis, candidates, new_name)


def prefixed_name(filename: str, mapping: MappingRow) -> str:
    target_prefix = f"{safe_component(mapping.objective)}_{safe_component(mapping.erl)}"
    stem, suffix = Path(filename).stem, Path(filename).suffix
    if stem.casefold().startswith(target_prefix.casefold() + "_"):
        return filename
    # Replace an existing recognizable Objective_ID + ERL prefix.
    old = re.match(
        r"(?i)^(?:[A-Z]{2}\.L2-)?3\.\d{1,2}\.\d{1,2}(?:\[[a-z]\])?"
        r"[_ -]+(?:(?:ERL[_ #.-]*)[A-Z0-9.-]+|E-[A-Z]{2}-\d+)[_ -]+(.+)$",
        stem,
    )
    base = old.group(1) if old else stem
    # Avoid repeating a target identifier already used as the original prefix.
    for token in (mapping.objective, mapping.erl):
        token_match = re.match(r"(?i)^" + re.escape(token) + r"[_ -]+(.+)$", base)
        if token_match:
            base = token_match.group(1)
    proposed = f"{target_prefix}_{base}{suffix}"
    if len(proposed) <= 240:
        return proposed
    name_tag = hashlib.sha256(proposed.encode("utf-8")).hexdigest()[:8]
    available = max(1, 240 - len(target_prefix) - len(suffix) - len(name_tag) - 4)
    return f"{target_prefix}_{base[:available]}__{name_tag}{suffix}"


def scan_paths(source_root: Path, mappings: list[MappingRow], selected: Iterable[Path] | None = None) -> list[FilePlan]:
    if selected is None:
        files = sorted((p for p in source_root.rglob("*") if p.is_file()), key=lambda p: str(p).casefold())
    else:
        files = sorted((p for p in selected if p.is_file()), key=lambda p: str(p).casefold())
    plans = []
    for path in files:
        if path.name.startswith("change_log_") or path.name in {"Thumbs.db", ".DS_Store"}:
            continue
        try:
            relative = path.relative_to(source_root)
        except ValueError:
            relative = Path(path.name)
        plans.append(match_file(path, relative, mappings))
    return plans


def scan_sources(
    source_roots: Iterable[Path],
    mappings: list[MappingRow],
    selected_files: Iterable[Path] | None = None,
) -> list[FilePlan]:
    """Scan one or more folders plus optional individual files.

    A single source folder retains the original v1.0 output layout. Multiple
    source groups receive a top-level folder label to keep each package separate.
    """
    roots = list(dict.fromkeys(Path(root).resolve() for root in source_roots))
    files = list(dict.fromkeys(Path(path).resolve() for path in (selected_files or [])))
    group_count = len(roots) + (1 if files else 0)
    plans: list[FilePlan] = []
    used_labels: dict[str, int] = {}

    def group_label(preferred: str) -> str:
        safe = safe_component(preferred) or "Evidence"
        count = used_labels.get(safe.casefold(), 0) + 1
        used_labels[safe.casefold()] = count
        return safe if count == 1 else f"{safe}__source{count:02d}"

    for root in roots:
        label = group_label(root.name)
        for plan in scan_paths(root, mappings):
            if group_count > 1:
                plan.relative_path = Path(label) / plan.relative_path
            plans.append(plan)
    if files:
        label = group_label("Selected_Files")
        common_names: dict[str, int] = {}
        for path in sorted((p for p in files if p.is_file()), key=lambda p: str(p).casefold()):
            name_count = common_names.get(path.name.casefold(), 0) + 1
            common_names[path.name.casefold()] = name_count
            relative = Path(path.name) if name_count == 1 else Path(f"source{name_count:02d}_{path.name}")
            if group_count > 1:
                relative = Path(label) / relative
            plans.append(match_file(path, relative, mappings))
    return sorted(plans, key=lambda plan: str(plan.relative_path).casefold())


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def unique_destination(path: Path) -> Path:
    if not path.exists():
        return path
    for number in range(1, 10000):
        candidate = path.with_name(f"{path.stem}__dup{number:02d}{path.suffix}")
        if not candidate.exists():
            return candidate
    raise RuntimeError(f"Could not create a unique output name for {path.name}")


LOG_FIELDS = [
    "timestamp_utc", "status", "match_basis", "objective_id", "erl_number",
    "mapping_row", "source_relative_path", "output_relative_path", "original_filename",
    "new_filename", "size_bytes", "source_sha256", "output_sha256", "hash_verified", "note",
]


def execute_plans(plans: list[FilePlan], source_root: Path, output_root: Path) -> tuple[Path, dict[str, int]]:
    output_root.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    log_path = output_root / f"change_log_{timestamp}.csv"
    counts: dict[str, int] = {}
    with log_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=LOG_FIELDS)
        writer.writeheader()
        for plan in plans:
            now = datetime.now(timezone.utc).isoformat(timespec="seconds").replace("+00:00", "Z")
            row = {
                "timestamp_utc": now,
                "status": plan.status,
                "match_basis": plan.match_basis,
                "objective_id": plan.mapping.objective if plan.mapping else "",
                "erl_number": plan.mapping.erl if plan.mapping else "",
                "mapping_row": plan.mapping.row_number if plan.mapping else "",
                "source_relative_path": str(plan.relative_path),
                "output_relative_path": "",
                "original_filename": plan.source.name,
                "new_filename": plan.new_name,
                "size_bytes": plan.source.stat().st_size,
                "source_sha256": "",
                "output_sha256": "",
                "hash_verified": "",
                "note": plan.note,
            }
            if plan.status in {"Ready", "Already aligned"} and plan.mapping:
                destination_dir = output_root / plan.relative_path.parent
                destination_dir.mkdir(parents=True, exist_ok=True)
                destination = unique_destination(destination_dir / plan.new_name)
                source_hash = sha256(plan.source)
                shutil.copy2(plan.source, destination)
                output_hash = sha256(destination)
                verified = source_hash == output_hash
                final_status = "Copied" if verified else "Hash mismatch"
                row.update({
                    "status": final_status,
                    "output_relative_path": str(destination.relative_to(output_root)),
                    "new_filename": destination.name,
                    "source_sha256": source_hash,
                    "output_sha256": output_hash,
                    "hash_verified": str(verified).upper(),
                })
                counts[final_status] = counts.get(final_status, 0) + 1
            else:
                counts[plan.status] = counts.get(plan.status, 0) + 1
            writer.writerow(row)
    return log_path, counts


def run_gui() -> None:
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk

    class App(tk.Tk):
        def __init__(self) -> None:
            super().__init__()
            self.title("CMMC Evidence Prefix Manager")
            self.geometry("1280x760")
            self.minsize(980, 620)
            self.mapping_path = tk.StringVar()
            self.source_path = tk.StringVar()
            self.output_path = tk.StringVar()
            self.status_text = tk.StringVar(value="Select a mapping file and evidence folder, then choose Preview.")
            self.mappings: list[MappingRow] = []
            self.plans: list[FilePlan] = []
            self.source_roots: list[Path] = []
            self.selected_files: list[Path] = []
            self._build()

        def _build(self) -> None:
            root = ttk.Frame(self, padding=12)
            root.pack(fill="both", expand=True)
            ttk.Label(root, text="CMMC Evidence Prefix Manager", font=("Segoe UI", 17, "bold")).pack(anchor="w")
            ttk.Label(root, text="Preview and copy evidence using ObjectiveID_ERL# filename prefixes.").pack(anchor="w", pady=(2, 12))
            form = ttk.Frame(root)
            form.pack(fill="x")
            self._picker(form, 0, "DRL mapping", self.mapping_path, self.pick_mapping)
            self._picker(form, 1, "Evidence source", self.source_path, self.pick_source, self.pick_files)
            self._picker(form, 2, "Output folder", self.output_path, self.pick_output)
            ttk.Button(form, text="Clear sources", command=self.clear_sources).grid(row=1, column=4, padx=(5, 0))
            actions = ttk.Frame(root)
            actions.pack(fill="x", pady=10)
            ttk.Button(actions, text="Preview", command=self.preview).pack(side="left")
            ttk.Button(actions, text="Apply selected mapping", command=self.override_selected).pack(side="left", padx=6)
            ttk.Button(actions, text="Skip selected", command=self.skip_selected).pack(side="left")
            self.mapping_combo = ttk.Combobox(actions, state="readonly", width=48)
            self.mapping_combo.pack(side="left", padx=6)
            self.run_button = ttk.Button(actions, text="Export renamed copies", command=self.export, state="disabled")
            self.run_button.pack(side="right")

            columns = ("status", "source", "objective", "erl", "new_name", "basis")
            self.tree = ttk.Treeview(root, columns=columns, show="headings", selectmode="extended")
            widths = {"status": 110, "source": 300, "objective": 150, "erl": 120, "new_name": 330, "basis": 180}
            labels = {"status": "Status", "source": "Source", "objective": "Objective ID", "erl": "ERL #", "new_name": "New filename", "basis": "Match basis"}
            for column in columns:
                self.tree.heading(column, text=labels[column])
                self.tree.column(column, width=widths[column], minwidth=80)
            scroll_y = ttk.Scrollbar(root, orient="vertical", command=self.tree.yview)
            scroll_x = ttk.Scrollbar(root, orient="horizontal", command=self.tree.xview)
            self.tree.configure(yscrollcommand=scroll_y.set, xscrollcommand=scroll_x.set)
            self.tree.pack(fill="both", expand=True, side="left")
            scroll_y.pack(fill="y", side="right")
            scroll_x.pack(fill="x", side="bottom")
            ttk.Label(self, textvariable=self.status_text, relief="sunken", anchor="w", padding=6).pack(fill="x", side="bottom")

        def _picker(self, parent, row, label, variable, command, file_command=None) -> None:
            ttk.Label(parent, text=label, width=16).grid(row=row, column=0, sticky="w", pady=3)
            ttk.Entry(parent, textvariable=variable).grid(row=row, column=1, sticky="ew", padx=5)
            ttk.Button(parent, text="Browse folder" if file_command else "Browse", command=command).grid(row=row, column=2)
            if file_command:
                ttk.Button(parent, text="Select files", command=file_command).grid(row=row, column=3, padx=(5, 0))
            parent.columnconfigure(1, weight=1)

        def pick_mapping(self) -> None:
            value = filedialog.askopenfilename(filetypes=[("Mapping files", "*.csv *.tsv *.xlsx *.xlsm"), ("All files", "*.*")])
            if value:
                self.mapping_path.set(value)

        def pick_source(self) -> None:
            value = filedialog.askdirectory()
            if value:
                root = Path(value).resolve()
                if root not in self.source_roots:
                    self.source_roots.append(root)
                self.update_source_display()
                if not self.output_path.get():
                    self.output_path.set(str(root.parent / "CMMC_Evidence_Renamed"))

        def pick_files(self) -> None:
            values = filedialog.askopenfilenames(title="Select evidence files")
            if values:
                for value in values:
                    path = Path(value).resolve()
                    if path not in self.selected_files:
                        self.selected_files.append(path)
                common = self.selected_files[0].parent
                self.update_source_display()
                if not self.output_path.get():
                    self.output_path.set(str(common.parent / "CMMC_Evidence_Renamed"))
                self.status_text.set(f"Selected {len(self.selected_files)} individual evidence file(s). Choose Preview.")

        def update_source_display(self) -> None:
            parts = [str(root) for root in self.source_roots]
            if self.selected_files:
                parts.append(f"[{len(self.selected_files)} individual file(s)]")
            self.source_path.set(" ; ".join(parts))

        def clear_sources(self) -> None:
            self.source_roots.clear()
            self.selected_files.clear()
            self.source_path.set("")
            self.plans.clear()
            self.refresh()
            self.run_button.config(state="disabled")

        def pick_output(self) -> None:
            value = filedialog.askdirectory(mustexist=False)
            if value:
                self.output_path.set(value)

        def preview(self) -> None:
            try:
                mapping_file = Path(self.mapping_path.get())
                if not mapping_file.is_file() or not (self.source_roots or self.selected_files):
                    raise ValueError("Select a valid DRL mapping file and at least one evidence folder or file.")
                output = Path(self.output_path.get()).resolve() if self.output_path.get() else None
                for source in self.source_roots:
                    if output and (output == source or source in output.parents):
                        raise ValueError("The output folder must be outside every source evidence folder.")
                self.mappings = load_mapping(mapping_file)
                self.plans = scan_sources(self.source_roots, self.mappings, self.selected_files)
                self.mapping_combo["values"] = [m.label for m in self.mappings]
                if self.mappings:
                    self.mapping_combo.current(0)
                self.refresh()
                self.run_button.config(state="normal" if self.plans else "disabled")
            except Exception as exc:
                messagebox.showerror("Preview failed", str(exc))

        def refresh(self) -> None:
            self.tree.delete(*self.tree.get_children())
            counts: dict[str, int] = {}
            for index, plan in enumerate(self.plans):
                mapping = plan.mapping
                self.tree.insert("", "end", iid=str(index), values=(
                    plan.status, str(plan.relative_path), mapping.objective if mapping else "",
                    mapping.erl if mapping else "", plan.new_name, plan.match_basis,
                ))
                counts[plan.status] = counts.get(plan.status, 0) + 1
            summary = ", ".join(f"{key}: {value}" for key, value in sorted(counts.items()))
            self.status_text.set(f"Mapping rows: {len(self.mappings)} | Files: {len(self.plans)} | {summary}")

        def override_selected(self) -> None:
            selected = self.tree.selection()
            index = self.mapping_combo.current()
            if not selected or index < 0:
                messagebox.showinfo("Manual mapping", "Select one or more files and a mapping row.")
                return
            mapping = self.mappings[index]
            for item in selected:
                plan = self.plans[int(item)]
                plan.mapping = mapping
                plan.status = "Ready"
                plan.match_basis = "Manual override"
                plan.new_name = prefixed_name(plan.source.name, mapping)
            self.refresh()

        def skip_selected(self) -> None:
            for item in self.tree.selection():
                plan = self.plans[int(item)]
                plan.status = "Skipped"
                plan.note = "Skipped by reviewer"
            self.refresh()

        def export(self) -> None:
            unresolved = sum(p.status in {"Unmatched", "Ambiguous"} for p in self.plans)
            if unresolved and not messagebox.askyesno(
                "Unresolved files",
                f"{unresolved} file(s) are unmatched or ambiguous. They will be logged but not copied. Continue?",
            ):
                return
            output = Path(self.output_path.get())
            if not str(output):
                messagebox.showerror("Output required", "Select an output folder.")
                return
            self.run_button.config(state="disabled")
            self.status_text.set("Export in progress. Hashing large files may take a moment.")

            def worker() -> None:
                try:
                    log_path, counts = execute_plans(self.plans, Path("."), output)
                    summary = ", ".join(f"{k}: {v}" for k, v in sorted(counts.items()))
                    self.after(0, lambda: messagebox.showinfo("Export complete", f"{summary}\n\nChange log:\n{log_path}"))
                    self.after(0, lambda: self.status_text.set(f"Export complete | {summary}"))
                except Exception as exc:
                    self.after(0, lambda: messagebox.showerror("Export failed", str(exc)))
                    self.after(0, lambda: self.status_text.set("Export failed. No source files were modified."))
                finally:
                    self.after(0, lambda: self.run_button.config(state="normal"))
            threading.Thread(target=worker, daemon=True).start()

    App().mainloop()


def cli_main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Prefix CMMC evidence filenames from a DRL mapping.")
    parser.add_argument("--mapping", type=Path, help="CSV, TSV, XLSX, or XLSM mapping file")
    parser.add_argument("--source", type=Path, action="append", help="Source evidence folder; repeat for multiple folders")
    parser.add_argument("--output", type=Path, help="Output folder for renamed copies")
    parser.add_argument("--execute", action="store_true", help="Copy Ready files after previewing results")
    parser.add_argument("--gui", action="store_true", help="Open the desktop GUI")
    args = parser.parse_args(argv)
    if args.gui or not any((args.mapping, args.source, args.output)):
        run_gui()
        return 0
    if not all((args.mapping, args.source, args.output)):
        parser.error("--mapping, --source, and --output are required together")
    mappings = load_mapping(args.mapping)
    plans = scan_sources(args.source, mappings)
    for plan in plans:
        objective = plan.mapping.objective if plan.mapping else ""
        erl = plan.mapping.erl if plan.mapping else ""
        print(f"{plan.status:15} | {objective:18} | {erl:12} | {plan.relative_path}")
    counts: dict[str, int] = {}
    for plan in plans:
        counts[plan.status] = counts.get(plan.status, 0) + 1
    print("\nPreview:", ", ".join(f"{key}={value}" for key, value in sorted(counts.items())))
    if args.execute:
        log_path, result_counts = execute_plans(plans, Path("."), args.output)
        print("Export:", ", ".join(f"{key}={value}" for key, value in sorted(result_counts.items())))
        print("Change log:", log_path)
    else:
        print("No files copied. Add --execute after reviewing the preview.")
    return 0


if __name__ == "__main__":
    raise SystemExit(cli_main())
