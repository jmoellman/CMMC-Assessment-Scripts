"""
Report output for the CMMC Pre-Assessment Package Checker.
  - HTML report (view in browser, print, or "Save as PDF" from the print dialog)
  - Excel issue log
  - CSV issue log
"""
from __future__ import annotations

import csv
import html
import os

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from engine import APP_NAME, APP_VERSION, FILE_LABELS, SEVERITIES, Result

E = html.escape

CSS = """
:root{--orange:#F7941D;--charcoal:#53555F;--dark:#292C35;--darktint:#3D404C;--light:#F4F4F5;--mid:#9A9BA3;--tint:#FEF0DC;--line:#E5E5E7}
*{box-sizing:border-box}
body{margin:0;background:#fff;color:var(--charcoal);font-family:'Proxima Nova','Inter','Open Sans',Arial,sans-serif;font-size:13px;line-height:1.4}
header{background:var(--dark);color:#fff;padding:18px 28px;border-bottom:4px solid var(--orange)}
header .brand{font-family:'Noway','Montserrat','Gill Sans',sans-serif;font-weight:700;font-size:13px;letter-spacing:.5px;color:var(--orange)}
header h1{font-family:'Noway','Montserrat','Gill Sans',sans-serif;font-size:22px;margin:4px 0 2px;font-weight:700}
header .sub{color:#d9dade;font-size:12px}
main{padding:20px 28px 40px}
h2{font-family:'Noway','Montserrat','Gill Sans',sans-serif;color:var(--dark);font-size:17px;margin:26px 0 8px;padding-bottom:4px;border-bottom:2px solid var(--orange)}
h3{font-size:14px;color:var(--charcoal);margin:18px 0 6px}
.cards{display:flex;gap:12px;flex-wrap:wrap;margin:10px 0}
.card{border:1px solid var(--line);border-top:4px solid var(--mid);border-radius:8px;padding:10px 16px;min-width:130px;background:#fff}
.card .n{font-family:'Noway','Montserrat',sans-serif;font-size:28px;font-weight:700;color:var(--dark)}
.card .l{font-size:12px;color:var(--charcoal)}
.card.err{border-top-color:var(--orange)} .card.err .n{color:var(--orange)}
.card.warn{border-top-color:var(--darktint)}
table{width:100%;border-collapse:collapse;margin:6px 0 14px}
th{background:var(--dark);color:#fff;text-align:left;font-weight:600;padding:6px 8px;font-size:12px}
td{border-bottom:1px solid var(--line);padding:6px 8px;vertical-align:top;font-size:12px}
tr:nth-child(even) td{background:#fafafa}
.sev{display:inline-block;padding:1px 8px;border-radius:10px;font-size:11px;font-weight:700;white-space:nowrap}
.sev.Error{background:var(--orange);color:var(--dark)}
.sev.Warning{background:var(--tint);color:var(--dark);border:1px solid var(--orange)}
.sev.Info{background:var(--light);color:var(--charcoal);border:1px solid var(--mid)}
.own{font-size:11px;color:var(--mid)}
.val{font-family:Consolas,'Courier New',monospace;font-size:11px;color:var(--dark);word-break:break-word}
.fix{color:var(--dark)}
.meta td{border:none;padding:2px 8px 2px 0}
.callout{background:var(--tint);border-left:4px solid var(--orange);padding:8px 12px;margin:10px 0}
.small{font-size:11px;color:var(--mid)}
.ok{color:var(--charcoal);font-style:italic}
footer{color:var(--mid);font-size:11px;padding:10px 28px;border-top:1px solid var(--line)}
@media print{
  @page{size:landscape;margin:0.45in}
  header{-webkit-print-color-adjust:exact;print-color-adjust:exact}
  th,.sev,.card{-webkit-print-color-adjust:exact;print-color-adjust:exact}
  thead{display:table-header-group} tr{page-break-inside:avoid}
  .noprint{display:none}
  h2{page-break-after:avoid}
}
.noprint{margin:10px 0}
.noprint button{background:var(--orange);color:var(--dark);border:none;border-radius:4px;padding:7px 16px;font-weight:700;cursor:pointer}
"""


def html_report(res: Result, path: str, auto_print=False, owner_filter=None, severities=None):
    sev_ok = set(severities or SEVERITIES)
    issues = [i for i in res.sorted_issues() if i.severity in sev_ok and (not owner_filter or i.owner == owner_filter)]
    c = {s: sum(1 for i in issues if i.severity == s) for s in SEVERITIES}
    osc = ""
    for lab, val, _ in res.data.get("pa_fields", []):
        if lab == "OSC Name":
            osc = val
    osc = osc or res.data.get("sap_osc_name", "")
    out = []
    w = out.append
    w("<!doctype html><html lang='en'><head><meta charset='utf-8'>")
    w(f"<title>Pre-Assessment Check Report{(' - ' + E(osc)) if osc else ''}</title>")
    w(f"<meta name='viewport' content='width=device-width,initial-scale=1'><style>{CSS}</style></head><body>")
    w("<header><div class='brand'>IGNYTE ASSURANCE PLATFORM</div>")
    w(f"<h1>CMMC Level 2 pre-assessment package check</h1>")
    w(f"<div class='sub'>{E(osc) + ' | ' if osc else ''}Generated {E(res.generated)} | {APP_NAME} v{APP_VERSION}</div></header><main>")
    w("<div class='noprint'><button onclick='window.print()'>Print or save as PDF</button></div>")
    if owner_filter:
        w(f"<div class='callout'>Filtered view: items owned by <b>{E(owner_filter)}</b> only.</div>")

    # Summary
    w("<h2>Summary</h2><div class='cards'>")
    w(f"<div class='card err'><div class='n'>{c['Error']}</div><div class='l'>Errors (must fix)</div></div>")
    w(f"<div class='card warn'><div class='n'>{c['Warning']}</div><div class='l'>Warnings (review)</div></div>")
    w(f"<div class='card'><div class='n'>{c['Info']}</div><div class='l'>Info (confirm)</div></div>")
    osc_n = sum(1 for i in issues if i.owner == "OSC")
    ig_n = sum(1 for i in issues if i.owner == "Ignyte")
    w(f"<div class='card'><div class='n'>{osc_n}</div><div class='l'>OSC action items</div></div>")
    w(f"<div class='card'><div class='n'>{ig_n}</div><div class='l'>Ignyte action items</div></div></div>")

    w("<table><thead><tr><th>File</th><th>File name</th><th>Errors</th><th>Warnings</th><th>Info</th><th>SHA-256</th></tr></thead><tbody>")
    for key, lab in FILE_LABELS.items():
        f = res.files.get(key)
        if key != "x" and not f:
            w(f"<tr><td>{E(lab)}</td><td class='ok'>Not loaded</td><td></td><td></td><td></td><td></td></tr>")
            continue
        if key == "x" and len(res.files) < 2:
            continue
        cnt = {s: sum(1 for i in issues if i.file_key == key and i.severity == s) for s in SEVERITIES}
        w(f"<tr><td>{E(lab)}</td><td>{E(f['name']) if f else 'Consistency across loaded files'}</td>"
          f"<td>{cnt['Error']}</td><td>{cnt['Warning']}</td><td>{cnt['Info']}</td>"
          f"<td class='val'>{E(f['sha256']) if f else ''}</td></tr>")
    w("</tbody></table>")
    w("<p class='small'>Severity: Error = blocks eMASS import or a required SAP/template item is missing. Warning = likely problem to review. "
      "Info = confirm with the OSC or Ignyte team. Owner shows who is expected to act (per SAP section 5.1, Ignyte completes the C3PAO date and standard rows).</p>")

    # Issues by file
    for key, lab in FILE_LABELS.items():
        rows = [i for i in issues if i.file_key == key]
        if key not in res.files and key != "x":
            continue
        if key == "x" and len(res.files) < 2:
            continue
        w(f"<h2>{E(lab)}</h2>")
        if not rows:
            w("<p class='ok'>No issues found.</p>")
            continue
        w("<table><thead><tr><th style='width:7%'>Severity</th><th style='width:14%'>Location</th><th style='width:13%'>Field</th>"
          "<th style='width:28%'>Issue</th><th style='width:16%'>Found</th><th style='width:22%'>How to resolve</th></tr></thead><tbody>")
        for i in rows:
            exp = f"<div class='small'>Expected: {E(i.expected)}</div>" if i.expected else ""
            af = "<div class='small'>Auto-fix available in corrected copy</div>" if i.autofix else ""
            w(f"<tr><td><span class='sev {i.severity}'>{i.severity}</span><div class='own'>{E(i.owner)}</div></td>"
              f"<td>{E(i.location)}</td><td>{E(i.field)}</td><td>{E(i.message)}</td>"
              f"<td class='val'>{E(i.value)}{exp}</td><td class='fix'>{E(i.fix)}{af}</td></tr>")
        w("</tbody></table>")

    # Captured data
    d = res.data
    w("<h2>Captured data for eMASS</h2>")
    w("<p class='small'>Values as read from the files. Use this to confirm what will carry into the eMASS Pre-Assessment and ESP records.</p>")
    if d.get("pa_fields"):
        w(f"<h3>Pre-Assessment form (template version {E(str(d.get('pa_template_version', '')))})</h3>")
        w("<table><thead><tr><th style='width:30%'>Field</th><th>Value read</th><th style='width:10%'>Input cell</th></tr></thead><tbody>")
        for lab, val, cell in d["pa_fields"]:
            w(f"<tr><td>{E(lab)}</td><td class='val'>{E(val) if val else '<span class=ok>blank</span>'}</td><td>{E(cell)}</td></tr>")
        w("</tbody></table>")
    if "esps" in d:
        w("<h3>External Service Providers (Asset Inventory, Consolidated ESP tab)</h3>")
        if d["esps"]:
            w("<table><thead><tr><th>ESP name</th><th>POC</th><th>Phone</th><th>Email</th><th>Description</th><th>CMMC status</th><th>Sector</th></tr></thead><tbody>")
            for e in d["esps"]:
                w(f"<tr><td>{E(e.get('name',''))}</td><td>{E(e.get('first',''))} {E(e.get('last',''))}</td><td>{E(e.get('phone',''))}</td>"
                  f"<td>{E(e.get('email',''))}</td><td>{E(e.get('desc',''))}</td><td>{E(e.get('status',''))}</td><td>{E(e.get('sector',''))}</td></tr>")
            w("</tbody></table>")
        else:
            w("<p class='ok'>No ESP records (sample record excluded).</p>")
    if "assets" in d:
        cats = {}
        for a in d["assets"]:
            cats[a.get("cat") or "(blank)"] = cats.get(a.get("cat") or "(blank)", 0) + 1
        w("<h3>Asset inventory by category</h3>")
        if cats:
            w("<table><thead><tr><th>Asset category</th><th>Rows</th></tr></thead><tbody>")
            for k, v in sorted(cats.items()):
                w(f"<tr><td>{E(k)}</td><td>{v}</td></tr>")
            w("</tbody></table>")
        else:
            w("<p class='ok'>No assets listed.</p>")
    if "team" in d:
        w("<h3>OSC Assessment Team (DRL)</h3>")
        if d["team"]:
            w("<table><thead><tr><th>Name</th><th>Email</th><th>Title</th><th>Responsibility</th></tr></thead><tbody>")
            for t in d["team"]:
                w(f"<tr><td>{E(t.get('name',''))}</td><td>{E(t.get('email',''))}</td><td>{E(t.get('title',''))}</td><td>{E(t.get('resp',''))}</td></tr>")
            w("</tbody></table>")
        else:
            w("<p class='ok'>No team members listed.</p>")
    if "facilities" in d:
        w("<h3>Physical facilities (DRL)</h3>")
        if d["facilities"]:
            w("<table><thead><tr><th>Location</th><th>Address</th><th>Description</th><th>Ownership</th><th>CAGE</th><th>Cleared</th><th>Onsite</th></tr></thead><tbody>")
            for f in d["facilities"]:
                w("<tr>" + "".join(f"<td>{E(f.get(k,''))}</td>" for k in ("name", "addr", "desc", "own", "cage", "cleared", "onsite")) + "</tr>")
            w("</tbody></table>")
        else:
            w("<p class='ok'>No facilities listed.</p>")
    if d.get("sap_yn"):
        w("<h3>SAP affirmations</h3><table><thead><tr><th>Item</th><th style='width:15%'>Answer</th></tr></thead><tbody>")
        for lab, a in d["sap_yn"].items():
            txt = "Yes" if a is True else "No" if a is False else "Not answered"
            w(f"<tr><td>{E(lab)}</td><td>{txt}</td></tr>")
        w("</tbody></table>")

    w("</main><footer>Ignyte Assurance Platform | Pre-assessment package check. Results reflect the files as loaded and do not represent an assessment determination.</footer>")
    if auto_print:
        w("<script>window.addEventListener('load',()=>setTimeout(()=>window.print(),400));</script>")
    w("</body></html>")
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("".join(out))
    return path


COLS = ["#", "Severity", "Owner", "File", "Location", "Field", "Issue", "Found", "Expected", "How to resolve", "Auto-fix", "Rule"]


def _rows(res):
    for n, i in enumerate(res.sorted_issues(), 1):
        yield [n, i.severity, i.owner, FILE_LABELS.get(i.file_key, i.file_key), i.location, i.field, i.message,
               i.value, i.expected, i.fix, "Yes" if i.autofix else "", i.rule]


def excel_report(res: Result, path: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Issues"
    navy = PatternFill("solid", fgColor="292C35")
    ws.append(COLS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF", name="Calibri")
        c.fill = navy
        c.alignment = Alignment(vertical="center", wrap_text=True)
    fills = {"Error": PatternFill("solid", fgColor="F7941D"), "Warning": PatternFill("solid", fgColor="FEF0DC"),
             "Info": PatternFill("solid", fgColor="F4F4F5")}
    for r in _rows(res):
        ws.append(r)
        ws.cell(ws.max_row, 2).fill = fills.get(r[1], PatternFill())
    widths = [5, 10, 9, 22, 28, 26, 60, 34, 28, 50, 9, 14]
    for i, wdt in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = wdt
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    s = wb.create_sheet("Files")
    s.append(["File", "File name", "SHA-256", "Generated"])
    for k, f in res.files.items():
        s.append([FILE_LABELS[k], f["name"], f["sha256"], res.generated])
    for c in s[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = navy
    for col, wdt in zip("ABCD", (26, 60, 70, 20)):
        s.column_dimensions[col].width = wdt

    if res.data.get("pa_fields"):
        p = wb.create_sheet("Pre-Assessment Data")
        p.append(["Field", "Value read", "Input cell"])
        for row in res.data["pa_fields"]:
            p.append(list(row))
        for c in p[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = navy
        p.column_dimensions["A"].width = 38
        p.column_dimensions["B"].width = 60
        p.column_dimensions["C"].width = 12
    wb.save(path)
    return path


def csv_report(res: Result, path: str):
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(COLS)
        for r in _rows(res):
            w.writerow(r)
    return path
