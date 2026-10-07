"""Validation engine for CMMC Level 2 eMASS Assessment Results workbooks (template v3.x).

Loads the workbook read-only, keeps a working copy of cell values, runs eMASS template
and assessor-guideline checks, and tracks accepted edits. Writing is done by xlsx_patch so
the template's dropdown validations, formulas, and formatting stay intact.
"""
from __future__ import annotations

import datetime as dt
import json
import os
import re
import warnings
from dataclasses import dataclass, field

import openpyxl
from openpyxl.utils import get_column_letter

from textcheck import APP_DIR, TextChecker, load_rules

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

S_ASM, S_REQ, S_OBJ, S_SSP = "Assessment", "Requirements", "Requirement Objectives", "OSC SSP(s)"
DATA_SHEETS = [S_ASM, S_REQ, S_OBJ, S_SSP]
ALL_SHEETS = DATA_SHEETS + ["Example", "Instructions", "Glossary", "Version History", "Lookup Values"]
DOMAINS = ["AC", "AT", "AU", "CA", "CM", "IA", "IR", "MA", "MP", "PE", "PS", "RA", "SC", "SI"]

# Requirement Objectives columns
OC = dict(req=1, obj=2, desc=3, artifacts=4, interviews=5, examine=6, test=7, overall=8, time=9, std=10,
          inherited=11, esp=12, score=13, date=14, by=15, findings=16)
OBJ_FIELD = {6: "Examine", 7: "Test", 8: "Overall Comments", 16: "Findings", 12: "ESP Name", 15: "Assessed By"}
EDITABLE = {(S_OBJ, c) for c in (6, 7, 8, 12, 16)} | {(S_ASM, 4)}
PROTECTED_NOTE = ("Protected field. Requirement Number, Objective Number, Artifacts, Interviews, Time to Assess, Score, "
                  "Findings status, and lookup-controlled values are not modified during this pass; flagged for assessor review.")

SCORE_VALUES = ["MET", "NOT MET", "NOT APPLICABLE"]
SCORE_VALUES_SC31311 = ["MET", "NOT MET -3", "NOT MET -5", "NOT APPLICABLE"]
INHERITED_VALUES = ["None", "Partial", "Full"]
STD_ACCEPT_VALUES = ["DIBCAC High", "FedRAMP Moderate", "FedRAMP High"]
POAM_VALUES = ["Yes", "No", "N/A"]
DATE_RX = re.compile(r"^\d{2}-(JAN|FEB|MAR|APR|MAY|JUN|JUL|AUG|SEP|OCT|NOV|DEC)-\d{4}$", re.I)

GLOSS = "eMASS Template Data Element Glossary"
INSTR = "eMASS Template Instructions"
GUIDE = "Assessor reporting guideline"

FIND_RX = re.compile(r"^(\s*)(NOT\s+MET(?:\s*-\s*[35])?|MET|N\s*/\s*A|NOT\s+APPLICABLE|NA)\b(\s*[:\-–—.,]?\s*)", re.I)


def norm_status(tok: str) -> str:
    t = re.sub(r"\s+", " ", tok.upper()).replace(" / ", "/")
    if t.startswith("NOT MET"):
        return "NOT MET"
    if t in ("N/A", "NA", "NOT APPLICABLE", "N /A", "N/ A"):
        return "NOT APPLICABLE"
    return "MET"


@dataclass
class Issue:
    uid: int
    severity: str
    category: str
    sheet: str
    row: int
    col: int
    field: str
    rule: str
    message: str
    why: str
    current: str = ""
    proposed: str | None = None
    safe: bool = False
    objective: str = ""
    span: tuple = (0, 0)
    word: str | None = None
    status: str = "Open"   # Open, Fixed, Ignored

    @property
    def col_letter(self):
        return get_column_letter(self.col) if self.col else ""

    @property
    def cell(self):
        return f"{self.col_letter}{self.row}" if self.col else (str(self.row) if self.row else "")

    @property
    def location(self):
        return f"{self.sheet}!{self.cell}" if self.cell else self.sheet

    def key(self):
        snippet = self.current[self.span[0]:self.span[1]] if self.span != (0, 0) else ""
        return (self.rule, self.sheet, self.cell, snippet, self.message if not snippet else "")


def as_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def parse_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = as_text(v).strip()
    if DATE_RX.match(s):
        try:
            return dt.datetime.strptime(s.title(), "%d-%b-%Y").date()
        except ValueError:
            return None
    return None


class Results:
    def __init__(self, path: str):
        self.path = path
        self.wb = openpyxl.load_workbook(path, data_only=False)
        with open(os.path.join(APP_DIR, "data", "template_reference.json"), encoding="utf-8") as fh:
            self.ref = json.load(fh)
        self.orig: dict = {}
        self.work: dict = {}
        for name in DATA_SHEETS:
            if name in self.wb.sheetnames:
                ws = self.wb[name]
                for row in ws.iter_rows():
                    for c in row:
                        if c.value is not None:
                            self.orig[(name, c.row, c.column)] = c.value
        self.work = dict(self.orig)
        self.changes: dict = {}     # (sheet,row,col) -> new text
        self.change_log: list = []  # dicts: location, rule, before, after
        # objective row map
        self.obj_rows = []
        if S_OBJ in self.wb.sheetnames:
            ws = self.wb[S_OBJ]
            cur = None
            for r in range(6, ws.max_row + 1):
                a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
                if a:
                    cur = str(a).strip()
                if b or a:
                    self.obj_rows.append((r, cur, str(b).strip() if b else ""))
        self.template_version = self._detect_version()

    def _detect_version(self):
        for name in DATA_SHEETS:
            if name not in self.wb.sheetnames:
                continue
            ws = self.wb[name]
            for c in ws[3]:
                if c.value and "Template Version" in str(c.value):
                    m = re.search(r"Template Version\s*([\d.]+)", str(c.value))
                    if m:
                        return m.group(1)
        return "unknown"

    def get(self, sheet, row, col):
        return self.work.get((sheet, row, col))

    def set_text(self, sheet, row, col, text, rule="Manual edit"):
        before = as_text(self.work.get((sheet, row, col)))
        if before == text:
            return
        self.work[(sheet, row, col)] = text
        self.changes[(sheet, row, col)] = text
        self.change_log.append({"location": f"{sheet}!{get_column_letter(col)}{row}", "rule": rule, "before": before, "after": text,
                                "time": dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")})

    def revert(self, sheet, row, col):
        o = self.orig.get((sheet, row, col))
        if (sheet, row, col) in self.changes:
            before = as_text(self.work.get((sheet, row, col)))
            del self.changes[(sheet, row, col)]
            if o is None:
                self.work.pop((sheet, row, col), None)
            else:
                self.work[(sheet, row, col)] = o
            self.change_log.append({"location": f"{sheet}!{get_column_letter(col)}{row}", "rule": "Reverted", "before": before, "after": as_text(o),
                                    "time": dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")})


# --------------------------------------------------------------------------------------
class Scope:
    def __init__(self, domains=None, obj_from=None, obj_to=None, assessment=True, requirements=True, ssp=True):
        self.domains = set(domains or DOMAINS)
        self.obj_from, self.obj_to = obj_from or None, obj_to or None
        self.assessment, self.requirements, self.ssp = assessment, requirements, ssp

    def describe(self):
        d = "All domains" if self.domains == set(DOMAINS) else ", ".join(sorted(self.domains))
        r = f"; objectives {self.obj_from or 'first'} to {self.obj_to or 'last'}" if (self.obj_from or self.obj_to) else ""
        extra = [n for n, on in (("Assessment", self.assessment), ("Requirements", self.requirements), ("OSC SSP(s)", self.ssp)) if on]
        return f"{d}{r}; sheets: Requirement Objectives" + ("".join(", " + x for x in extra))


class Engine:
    def __init__(self, results: Results, opts: dict | None = None, checker: TextChecker | None = None):
        self.res = results
        self.opts = opts or {}
        self.rules = load_rules()
        self.tc = checker or TextChecker(self.rules)
        self.issues: list[Issue] = []
        self.ignored: set = set()
        self._uid = 0
        self.scope = Scope()

    # ------------------------------------------------------------------
    def _new(self, sev, cat, sheet, row, col, fieldname, rule, msg, why, current="", proposed=None, safe=False, objective="", span=(0, 0), word=None):
        self._uid += 1
        iss = Issue(self._uid, sev, cat, sheet, row, col, fieldname, rule, msg, why, current, proposed, safe, objective, span, word)
        if iss.key() in self.ignored:
            iss.status = "Ignored"
        return iss

    def in_scope_rows(self):
        rows = []
        idx = {o: i for i, (_r, _q, o) in enumerate(self.res.obj_rows)}
        lo = idx.get(self.scope.obj_from, 0) if self.scope.obj_from else 0
        hi = idx.get(self.scope.obj_to, len(self.res.obj_rows) - 1) if self.scope.obj_to else len(self.res.obj_rows) - 1
        for i, (r, req, obj) in enumerate(self.res.obj_rows):
            dom = (req or obj or "")[:2]
            if dom in self.scope.domains and lo <= i <= hi:
                rows.append((r, req, obj))
        return rows

    # ------------------------------------------------------------------
    def run(self, scope: Scope | None = None, progress=None):
        if scope:
            self.scope = scope
        self.issues = []
        out = self.issues
        out += self.check_structure()
        self.asm_dates = self._asm_dates()
        rows = self.in_scope_rows()
        if self.scope.assessment and S_ASM in self.res.wb.sheetnames:
            out += self.check_assessment()
        if self.scope.ssp and S_SSP in self.res.wb.sheetnames:
            out += self.check_ssp()
        if self.scope.requirements and S_REQ in self.res.wb.sheetnames:
            out += self.check_requirements({req for _r, req, _o in rows})
        n = len(rows)
        for i, (r, req, obj) in enumerate(rows):
            out += self.check_objective_row(r, req, obj)
            if progress:
                progress(i + 1, n)
        self.issues.sort(key=lambda x: (DATA_SHEETS.index(x.sheet) if x.sheet in DATA_SHEETS else -1, x.row, x.col, {"Error": 0, "Warning": 1, "Info": 2}[x.severity]))
        return self.issues

    def _asm_dates(self):
        if S_ASM not in self.res.wb.sheetnames:
            return (None, None)
        ws = self.res.wb[S_ASM]
        start = end = None
        for r in range(6, 20):
            name = as_text(ws.cell(r, 1).value)
            if name == "Assessment Start Date":
                start = parse_date(self.res.get(S_ASM, r, 4))
            elif name == "Assessment End Date":
                end = parse_date(self.res.get(S_ASM, r, 4))
        return (start, end)

    # ------------------------------------------------------------------
    def check_structure(self):
        out, wb = [], self.res.wb
        for name in ALL_SHEETS:
            if name not in wb.sheetnames:
                sev = "Error" if name in DATA_SHEETS or name == "Lookup Values" else "Warning"
                out.append(self._new(sev, "Template", name, 0, 0, "Worksheet", "EM-SHEET", f'Worksheet "{name}" is missing or renamed.',
                                     f"{INSTR} #2: do not delete columns or sheets. eMASS may reject the import."))
        extra = [s for s in wb.sheetnames if s not in ALL_SHEETS]
        for s in extra:
            out.append(self._new("Warning", "Template", s, 0, 0, "Worksheet", "EM-SHEET-EXTRA", f'Unexpected worksheet "{s}".',
                                 f"{INSTR} #2: do not add sheets or columns; extra content may affect eMASS ingestion."))
        for name in DATA_SHEETS:
            if name not in wb.sheetnames:
                continue
            ws = wb[name]
            a1 = as_text(ws["A1"].value)
            if "CUI" not in a1.upper():
                out.append(self._new("Error", "Template", name, 1, 1, "Classification label", "EM-CLASS", "Classification label is missing or changed.",
                                     f'{INSTR} #2: do not delete the classification label ("***** CUI (When Filled In) *****").', a1))
            exp = self.res.ref["headers"].get(name, [])
            for ci, h in enumerate(exp, start=1):
                if h is None:
                    continue
                got = as_text(ws.cell(5, ci).value).strip()
                if got != h:
                    out.append(self._new("Error", "Template", name, 5, ci, "Header", "EM-HEADER", f'Column header changed: expected "{h}", found "{got}".',
                                         f"{INSTR} #2: do not delete, move, or add columns. eMASS maps data by template position.", got))
            if ws.max_column > len(exp) + 1 and name != S_OBJ:
                out.append(self._new("Warning", "Template", name, 5, ws.max_column, "Columns", "EM-COLUMNS", "Content found beyond the template's last column.",
                                     f"{INSTR} #2: do not add additional columns."))
        if self.res.template_version not in ("unknown",) and self.res.template_version != self.res.ref["template_version"]:
            out.append(self._new("Info", "Template", S_OBJ, 3, 0, "Template version", "EM-VERSION",
                                 f"Template version {self.res.template_version} differs from the reference version {self.res.ref['template_version']} built into this checker.",
                                 "Row and lookup checks assume template v3.9. Confirm the eMASS template in use."))
        # ID integrity
        if S_OBJ in wb.sheetnames:
            ws = wb[S_OBJ]
            for o in self.res.ref["objectives"]:
                r = o["row"]
                got_a = as_text(ws.cell(r, 1).value).strip()
                got_b = as_text(ws.cell(r, 2).value).strip()
                exp_a = o["req_cell"] or ""
                if got_b != (o["obj"] or ""):
                    out.append(self._new("Error", "Template", S_OBJ, r, 2, "Objective Number", "EM-OBJ-ID", f'Objective Number is "{got_b}", template expects "{o["obj"]}".',
                                         f"{INSTR} (Requirement Objectives #2): do not modify Objective Number. {PROTECTED_NOTE}", got_b, objective=o["obj"]))
                if got_a != exp_a:
                    out.append(self._new("Error" if exp_a else "Warning", "Template", S_OBJ, r, 1, "Requirement Number", "EM-REQ-ID",
                                         f'Requirement Number is "{got_a or "(blank)"}", template has "{exp_a or "(blank)"}".',
                                         f"{INSTR} (Requirement Objectives #2): do not modify Requirement Number. {PROTECTED_NOTE}", got_a, objective=o["obj"]))
        if S_REQ in wb.sheetnames:
            ws = wb[S_REQ]
            for q in self.res.ref["requirements"]:
                got = as_text(ws.cell(q["row"], 1).value).strip()
                if got != q["req"]:
                    out.append(self._new("Error", "Template", S_REQ, q["row"], 1, "Requirement Number", "EM-REQ-ID", f'Requirement Number is "{got}", template expects "{q["req"]}".',
                                         f"{INSTR} (Requirements #2): do not modify Requirement Number. {PROTECTED_NOTE}", got))
        return out

    # ------------------------------------------------------------------
    def _limit(self, sheet, row, col, fname, text, limit, objective=""):
        if len(text) > limit:
            return [self._new("Error", "eMASS Limit", sheet, row, col, fname, "EM-LIMIT", f"{fname} is {len(text):,} characters; the limit is {limit:,}.",
                              f"{GLOSS}: {fname} free text limit is {limit:,} characters. Over-limit text may be rejected or truncated on import.", text, objective=objective)]
        return []

    def _date(self, sheet, row, col, fname, v, required, objective=""):
        out = []
        if v is None or as_text(v).strip() == "":
            if required:
                out.append(self._new("Error", "Required", sheet, row, col, fname, "EM-REQUIRED", f"{fname} is required and blank.",
                                     f"{INSTR} #3: eMASS will deny a template import if a required field is blank.", objective=objective))
            return out, None
        d = parse_date(v)
        if d is None:
            out.append(self._new("Error", "Format", sheet, row, col, fname, "EM-DATE", f'{fname} "{as_text(v)}" is not a valid date in DD-MMM-YYYY format.',
                                 f"{INSTR} #6: for date fields use the DD-MMM-YYYY format (for example 05-Oct-2026).", as_text(v), objective=objective))
        return out, d

    def _text_issues(self, sheet, row, col, fname, text, objective=""):
        out = []
        hits = self.tc.check(text, fname, {"spelling": self.opts.get("spelling", True), "languagetool": self.opts.get("languagetool", False)})
        for h in hits:
            proposed = h.apply(text) if h.replacement is not None else None
            if proposed is not None and fname == "Findings":
                # never alter the formal result token through a text fix
                a, b = FIND_RX.match(text), FIND_RX.match(proposed)
                if a and (not b or norm_status(a.group(2)) != norm_status(b.group(2))):
                    proposed = None
            if proposed is not None and (sheet, col) not in EDITABLE:
                proposed = None
            out.append(self._new(h.severity, h.category, sheet, row, col, fname, h.rule, h.message, h.why, text, proposed, h.safe and proposed is not None,
                                 objective, (h.start, h.end), h.word))
        return out

    # ------------------------------------------------------------------
    def check_assessment(self):
        out, ws, S = [], self.res.wb[S_ASM], S_ASM
        fields = {}
        for r in range(6, 6 + len(self.res.ref["assessment_fields"])):
            name = as_text(ws.cell(r, 1).value).strip()
            req = as_text(ws.cell(r, 2).value).strip().lower() == "required"
            fields[name] = (r, req)
        get = lambda n: fields.get(n, (None, False))  # noqa: E731
        start = end = None
        for name, (r, req) in fields.items():
            v = self.res.get(S, r, 4)
            t = as_text(v).strip()
            if "Date" in name:
                o, d = self._date(S, r, 4, name, v, req)
                out += o
                if "Start" in name:
                    start = d
                elif "End" in name:
                    end = d
                continue
            if req and not t:
                out.append(self._new("Error", "Required", S, r, 4, name, "EM-REQUIRED", f"{name} is required and blank.",
                                     f"{INSTR} #3 and #7: all required fields in the Assessment worksheet must be completed."))
                continue
            if not t:
                continue
            if name == "OSC Name":
                out += self._limit(S, r, 4, name, t, 100)
            elif "CMMC Professional Number" in name:
                out += self._limit(S, r, 4, name, t, 50)
                if not re.fullmatch(r"\d+", t):
                    digits = re.sub(r"\D", "", t)
                    out.append(self._new("Error", "Format", S, r, 4, name, "EM-CPN", f'CPN "{t}" must be the ID number only.',
                                         f'{INSTR} #10: enter the ID number only; for "CCP-50" enter "50".', t, digits or None, False))
            elif name == "Standards Acceptance":
                if t not in STD_ACCEPT_VALUES:
                    out.append(self._new("Error", "Lookup", S, r, 4, name, "EM-LOOKUP", f'"{t}" is not an accepted Standards Acceptance value.',
                                         f"{INSTR} #5 / {GLOSS}: accepted values are {', '.join(STD_ACCEPT_VALUES)}.", t))
            elif name == "C3PAO Executive Summary":
                out += self._limit(S, r, 4, name, t, 2000)
                out += self._text_issues(S, r, 4, "Executive Summary", as_text(v))
            elif name.startswith("Hash Value"):
                out += self._limit(S, r, 4, name, t, 100)
                hx = t.replace(" ", "")
                if not re.fullmatch(r"[0-9A-Fa-f]+", hx):
                    out.append(self._new("Warning", "Format", S, r, 4, name, "EM-HASH", "Hash value contains non-hexadecimal characters.",
                                         "A hash digest is a hexadecimal string; confirm it was copied completely from the hashing output.", t))
                algo = as_text(self.res.get(S, get("Hash Algorithm")[0] or 0, 4)).upper()
                exp = 64 if "256" in algo else 128 if "512" in algo else 96 if "384" in algo else 40 if re.search(r"SHA-?1\b", algo) else 32 if "MD5" in algo else None
                if exp and len(hx) != exp:
                    out.append(self._new("Warning", "Format", S, r, 4, name, "EM-HASH-LEN", f"Hash value is {len(hx)} characters; {algo.strip()} produces {exp}.",
                                         "The digest length does not match the stated hash algorithm. Confirm the value and algorithm.", t))
            elif name.startswith("Hashed Data List"):
                out += self._limit(S, r, 4, name, t, 2000)
            elif name == "Hash Algorithm":
                out += self._limit(S, r, 4, name, t, 100)
        if start and end and end < start:
            r = get("Assessment End Date")[0]
            out.append(self._new("Error", "Consistency", S, r, 4, "Assessment End Date", "EM-DATE-ORDER", "Assessment End Date is before the Start Date.",
                                 "The assessment period cannot end before it starts."))
        hr = get("Hash Date")[0]
        hd = parse_date(self.res.get(S, hr, 4)) if hr else None
        if hd and start and hd < start:
            out.append(self._new("Warning", "Consistency", S, hr, 4, "Hash Date", "EM-HASH-DATE", "Hash Date is before the Assessment Start Date.",
                                 "Artifacts are hashed after evidence collection; confirm the hash date."))
        return out

    # ------------------------------------------------------------------
    def check_ssp(self):
        out, S = [], S_SSP
        ws = self.res.wb[S]
        complete = 0
        for r in range(6, max(ws.max_row, 6) + 1):
            vals = [self.res.get(S, r, c) for c in (1, 2, 3)]
            if all(as_text(v).strip() == "" for v in vals):
                continue
            missing = [n for n, v in zip(("SSP Name", "SSP Version", "SSP Date"), vals) if as_text(v).strip() == ""]
            for n in missing:
                c = ("SSP Name", "SSP Version", "SSP Date").index(n) + 1
                out.append(self._new("Error", "Required", S, r, c, n, "EM-REQUIRED", f"{n} is blank on a listed SSP row.",
                                     f"{INSTR} #13 / {GLOSS}: SSP name, version, and date are required for each SSP."))
            out += self._limit(S, r, 1, "SSP Name", as_text(vals[0]), 250)
            out += self._limit(S, r, 2, "SSP Version", as_text(vals[1]), 250)
            if as_text(vals[2]).strip():
                o, _d = self._date(S, r, 3, "SSP Date", vals[2], True)
                out += o
            if not missing:
                complete += 1
            if isinstance(vals[1], float):
                out.append(self._new("Info", "Format", S, r, 2, "SSP Version", "EM-SSP-VER", f"SSP Version is stored as a number ({vals[1]}).",
                                     'Excel may drop trailing zeros (for example "2.10" becomes 2.1). Confirm the version matches the SSP.', as_text(vals[1])))
        if complete == 0:
            out.append(self._new("Error", "Required", S, 6, 1, "SSP Name", "EM-SSP-NONE", "No SSP is listed.",
                                 f"{INSTR} #13: the SSP name, version, and date associated with the assessment scope are required."))
        return out

    # ------------------------------------------------------------------
    def _objective_scores(self):
        by_req = {}
        for r, req, obj in self.res.obj_rows:
            by_req.setdefault(req, []).append((r, obj, as_text(self.res.get(S_OBJ, r, OC["score"])).strip()))
        return by_req

    def check_requirements(self, reqs_in_scope):
        out, S = [], S_REQ
        ws = self.res.wb[S]
        scores = self._objective_scores()
        for r in range(6, ws.max_row + 1):
            req = as_text(ws.cell(r, 1).value).strip()
            if not req or req not in reqs_in_scope:
                continue
            allowed = as_text(self.res.get(S, r, 5)).strip()
            v = as_text(self.res.get(S, r, 6)).strip()
            if not v:
                out.append(self._new("Error", "Required", S, r, 6, "Requirement in POA&M", "EM-REQUIRED", "Requirement in POA&M is required and blank.",
                                     f"{INSTR} (Requirements #5): Requirement in POA&M is required for each requirement.", objective=req))
                continue
            if v not in POAM_VALUES:
                out.append(self._new("Error", "Lookup", S, r, 6, "Requirement in POA&M", "EM-LOOKUP", f'"{v}" is not an accepted value.',
                                     f"{GLOSS}: accepted values are Yes, No, N/A (exact spelling and case). {PROTECTED_NOTE}", v, objective=req))
                continue
            objs = scores.get(req, [])
            sc = [s for _r, _o, s in objs]
            any_nm = any(s.upper().startswith("NOT MET") for s in sc)
            all_done = sc and all(s for s in sc)
            if v == "Yes" and allowed == "No":
                out.append(self._new("Error", "Consistency", S, r, 6, "Requirement in POA&M", "EM-POAM-NOTALLOWED", "Requirement in POA&M is Yes, but POA&M Allowed is No.",
                                     f"{INSTR} (Requirements #5): select Yes only when a POA&M is allowed for the requirement. {PROTECTED_NOTE}", v, objective=req))
            if any_nm and v == "N/A":
                out.append(self._new("Error", "Consistency", S, r, 6, "Requirement in POA&M", "EM-POAM-NA-NOTMET", "Requirement in POA&M is N/A, but at least one objective is NOT MET.",
                                     f"{INSTR} (Requirements #5): N/A applies only when the requirement was met. {PROTECTED_NOTE}", v, objective=req))
            if all_done and not any_nm and v == "Yes":
                out.append(self._new("Warning", "Consistency", S, r, 6, "Requirement in POA&M", "EM-POAM-YES-MET", "Requirement in POA&M is Yes, but no objective is NOT MET.",
                                     f"{INSTR} (Requirements #5): select N/A if the requirement was met. {PROTECTED_NOTE}", v, objective=req))
            if all_done and not any_nm and v == "No":
                out.append(self._new("Warning", "Consistency", S, r, 6, "Requirement in POA&M", "EM-POAM-NO-MET", "Requirement in POA&M is No, but all objectives are MET or NOT APPLICABLE.",
                                     f"{INSTR} (Requirements #5): select N/A if the requirement was met and therefore does not require a POA&M. {PROTECTED_NOTE}", v, objective=req))
        return out

    # ------------------------------------------------------------------
    def _list_field(self, r, col, fname, text, obj):
        out = []
        if not text.strip():
            out.append(self._new("Warning", "Completeness", S_OBJ, r, col, fname, "AG-ROW-BLANK", f"{fname} is blank.",
                                 f"{GUIDE}: no objective row should be left without {fname.lower()}; reuse entries from related objectives where applicable. ({GLOSS} lists this field as optional.)",
                                 text, objective=obj))
            return out
        out += self._limit(S_OBJ, r, col, fname, text, 4000, obj)
        parts = text.split(";")
        for p in parts:
            if len(p.strip()) > 400:
                out.append(self._new("Error", "eMASS Limit", S_OBJ, r, col, fname, "EM-LIMIT-ITEM", f'One {fname} entry is {len(p.strip())} characters (limit 400 per value): "{p.strip()[:60]}..."',
                                     f"{INSTR} (Requirement Objectives #5): each semicolon-separated value has a 400 character maximum.", text, objective=obj))
        if any(not p.strip() for p in parts[:-1]) or re.search(r";\s*;", text):
            out.append(self._new("Warning", "Format", S_OBJ, r, col, fname, "EM-LIST-EMPTY", f"{fname} contains an empty entry (double or leading semicolon).",
                                 f"{INSTR} (Requirement Objectives #5): use a semicolon to separate values. {PROTECTED_NOTE}", text, objective=obj))
        if "\n" in text.strip() or (";" not in text and ("," in text and len(text.split(",")) >= 3)):
            out.append(self._new("Warning", "Format", S_OBJ, r, col, fname, "EM-LIST-SEP", f"{fname} appears to use line breaks or commas instead of semicolons.",
                                 f"{INSTR} (Requirement Objectives #5): use a semicolon to separate values within each list. {PROTECTED_NOTE}", text, objective=obj))
        seen, dups = set(), set()
        for p in parts:
            k = p.strip().lower()
            if k and k in seen:
                dups.add(p.strip())
            seen.add(k)
        if dups:
            out.append(self._new("Info", "Format", S_OBJ, r, col, fname, "EM-LIST-DUP", f"Duplicate {fname} entries: {', '.join(sorted(dups))[:200]}.",
                                 f"Duplicate entries count toward the 4,000 character limit. {PROTECTED_NOTE}", text, objective=obj))
        if fname == "Interviews" and re.search(r"assigned certified assessor", text, re.I):
            out.append(self._new("Error", "Assessor Language", S_OBJ, r, col, fname, "AG-ACTOR-PROHIBITED", '"The assigned certified assessor" appears in Interviews.',
                                 f"{GUIDE}: this phrase is prohibited in eMASS deliverables. {PROTECTED_NOTE}", text, objective=obj))
        return out

    def check_objective_row(self, r, req, obj):
        out, S = [], S_OBJ
        g = lambda c: self.res.get(S, r, c)  # noqa: E731
        t = lambda c: as_text(g(c))  # noqa: E731
        is_sc = (req or "").startswith("SC.L2-3.13.11")
        if all(not t(c).strip() for c in range(4, 17)):
            return [self._new("Error", "Required", S, r, OC["score"], "Objective row", "EM-ROW-EMPTY",
                              "Objective row is not completed (Inherited, Time to Assess, Score, Date Assessed, Assessed By, and Findings are blank).",
                              f"{INSTR} #3 and (Requirement Objectives #3): required fields must be completed; eMASS will deny the import if a required field is blank. "
                              "If this objective is outside your assigned range, narrow the Scope.", objective=obj)]
        # Required lookups and fields
        score = t(OC["score"]).strip()
        allowed_scores = SCORE_VALUES_SC31311 if is_sc else SCORE_VALUES
        if not score:
            out.append(self._new("Error", "Required", S, r, OC["score"], "Score", "EM-REQUIRED", "Score is required and blank.",
                                 f"{INSTR} (Requirement Objectives #3): Score is required.", objective=obj))
        elif score not in allowed_scores:
            hint = " (case must match exactly; template v3.9 uses uppercase)" if score.upper() in allowed_scores else ""
            out.append(self._new("Error", "Lookup", S, r, OC["score"], "Score", "EM-LOOKUP", f'Score "{score}" is not an accepted value{hint}.',
                                 f"{INSTR} (Requirement Objectives #7{', #8' if is_sc else ''}): accepted values are {', '.join(allowed_scores)}. {PROTECTED_NOTE}", score, objective=obj))
        inh = t(OC["inherited"]).strip()
        esp = t(OC["esp"]).strip()
        if not inh:
            out.append(self._new("Error", "Required", S, r, OC["inherited"], "Inherited", "EM-REQUIRED", "Inherited is required and blank.",
                                 f"{INSTR} (Requirement Objectives #3, #4): Inherited is required (None, Partial, or Full).", objective=obj))
        elif inh not in INHERITED_VALUES:
            out.append(self._new("Error", "Lookup", S, r, OC["inherited"], "Inherited", "EM-LOOKUP", f'Inherited "{inh}" is not an accepted value.',
                                 f"{INSTR} (Requirement Objectives #4): accepted values are None, Partial, Full. {PROTECTED_NOTE}", inh, objective=obj))
        elif inh in ("Partial", "Full") and not esp:
            out.append(self._new("Warning", "Consistency", S, r, OC["esp"], "ESP Name", "EM-ESP-MISSING", f"Inherited is {inh} but no ESP name is entered.",
                                 f"{INSTR} (Requirement Objectives #6): if inherited, describe the dependency in 'If dependent on ESP, enter ESP Name'.", objective=obj))
        elif inh == "None" and esp:
            out.append(self._new("Warning", "Consistency", S, r, OC["inherited"], "Inherited", "EM-ESP-CONFLICT", "An ESP name is entered but Inherited is None.",
                                 f"Confirm whether the objective is inherited. {PROTECTED_NOTE}", inh, objective=obj))
        if esp:
            out += self._limit(S, r, OC["esp"], "ESP Name", esp, 2000, obj)
        std = t(OC["std"]).strip()
        if std and std not in STD_ACCEPT_VALUES:
            out.append(self._new("Error", "Lookup", S, r, OC["std"], "Standards Acceptance", "EM-LOOKUP", f'"{std}" is not an accepted Standards Acceptance value.',
                                 f"{GLOSS}: accepted values are {', '.join(STD_ACCEPT_VALUES)}. {PROTECTED_NOTE}", std, objective=obj))
        tm = t(OC["time"]).strip()
        if not tm:
            out.append(self._new("Error", "Required", S, r, OC["time"], "Time to Assess (Minutes)", "EM-REQUIRED", "Time to Assess is required and blank.",
                                 f"{INSTR} (Requirement Objectives #3): Time to Assess (Minutes) is required.", objective=obj))
        elif not re.fullmatch(r"\d+(\.\d+)?", tm) or float(tm) <= 0:
            out.append(self._new("Error", "Format", S, r, OC["time"], "Time to Assess (Minutes)", "EM-TIME", f'Time to Assess "{tm}" is not a positive number of minutes.',
                                 f"{GLOSS}: Time To Assess is numeric free text. {PROTECTED_NOTE}", tm, objective=obj))
        o, d = self._date(S, r, OC["date"], "Date Assessed", g(OC["date"]), True, obj)
        out += o
        s_e = getattr(self, "asm_dates", (None, None))
        if d and s_e[0] and s_e[1] and not (s_e[0] <= d <= s_e[1]):
            out.append(self._new("Warning", "Consistency", S, r, OC["date"], "Date Assessed", "EM-DATE-RANGE",
                                 f"Date Assessed {d:%d-%b-%Y} is outside the assessment period ({s_e[0]:%d-%b-%Y} to {s_e[1]:%d-%b-%Y}).",
                                 f"Confirm the date. {PROTECTED_NOTE}", as_text(g(OC['date'])), objective=obj))
        by = t(OC["by"]).strip()
        if not by:
            out.append(self._new("Error", "Required", S, r, OC["by"], "Assessed By", "EM-REQUIRED", "Assessed By is required and blank.",
                                 f"{INSTR} (Requirement Objectives #3): Assessed By is required.", objective=obj))
        else:
            out += self._limit(S, r, OC["by"], "Assessed By", by, 100, obj)
            for h in self.tc.check(by, "Assessed By"):
                out.append(self._new(h.severity, h.category, S, r, OC["by"], "Assessed By", h.rule, h.message, h.why + " " + PROTECTED_NOTE, by, None, False, obj, (h.start, h.end)))
        # Lists
        out += self._list_field(r, OC["artifacts"], "Artifacts", t(OC["artifacts"]), obj)
        out += self._list_field(r, OC["interviews"], "Interviews", t(OC["interviews"]), obj)
        # Narratives
        for col, fname in ((OC["examine"], "Examine"), (OC["test"], "Test"), (OC["overall"], "Overall Comments")):
            txt = t(col)
            if txt.strip():
                out += self._limit(S, r, col, fname, txt, 4000, obj)
                out += self._text_issues(S, r, col, fname, txt, obj)
        if esp:
            pass
        out += self.check_findings(r, obj, score)
        if score == "MET" and not t(OC["examine"]).strip() and not t(OC["overall"]).strip() and not t(OC["test"]).strip():
            out.append(self._new("Warning", "Determination Support", S, r, OC["examine"], "Examine", "AG-MET-EVIDENCE", "MET with no Examine, Test, or Overall Comments narrative.",
                                 f"{GUIDE}: for MET, the narrative states the accepted evidence and implementation supporting the objective.", objective=obj))
        return out

    def check_findings(self, r, obj, score):
        out, S, col = [], S_OBJ, OC["findings"]
        text = as_text(self.res.get(S, r, col))
        fcfg = self.rules["findings"]
        labels, sep = fcfg["labels"], fcfg.get("separator", ":")
        if not text.strip():
            out.append(self._new("Error", "Required", S, r, col, "Findings", "EM-REQUIRED", "Findings is required and blank.",
                                 f"{INSTR} (Requirement Objectives #3): Findings is required.", objective=obj))
            return out
        out += self._limit(S, r, col, "Findings", text, 4000, obj)
        score_n = "NOT MET" if score.upper().startswith("NOT MET") else score.upper()
        m = FIND_RX.match(text)
        why_open = f'{GUIDE}: Findings open with the formal narrative result (MET, NOT MET, or N/A) followed by objective-specific factual rationale, formatted as "{labels["MET"]}{sep} ..."'
        if not m:
            prop = None
            if score_n in labels:
                body = text.lstrip()
                prop = f"{labels[score_n]}{sep} {body}"
            out.append(self._new("Error", "Findings", S, r, col, "Findings", "AG-FIND-OPENER", "Findings do not open with the formal result (MET, NOT MET, or N/A).",
                                 why_open + (" Proposed fix uses the Score column value." if prop else " Score is blank or invalid, so no fix is proposed."),
                                 text, prop, False, obj, (0, min(len(text), 30))))
            body_start = 0
        else:
            status = norm_status(m.group(2))
            canon = f"{labels[status]}{sep} "
            raw = m.group(0)
            body_start = m.end()
            if score_n in labels and status != score_n:
                out.append(self._new("Error", "Findings", S, r, col, "Findings", "AG-FIND-CONFLICT", f'Findings open with "{m.group(2).strip()}" but Score is "{score}".',
                                     f"The narrative result conflicts with the Score lookup. {PROTECTED_NOTE}", text, None, False, obj, (m.start(2), m.end(2))))
            elif raw.lstrip() != canon and text[m.end():].strip():
                prop = text[:m.start(2)] + canon + text[m.end():]
                if m.start(2) > 0:
                    prop = prop.lstrip()
                out.append(self._new("Warning", "Findings", S, r, col, "Findings", "AG-FIND-FORMAT", f'Result label "{raw.strip()}" does not match the standard format "{canon.strip()}".',
                                     why_open + " The result itself is unchanged.", text, prop, False, obj, (m.start(2), m.end())))
            body = text[body_start:]
            low = body.lower()
            if status == "MET":
                for term in fcfg["met_weakness_terms"]:
                    mm = re.search(r"\b" + re.escape(term.strip()) + (r"\b" if not term.strip().endswith(("deficien",)) else ""), body, re.I)
                    if mm:
                        out.append(self._new("Warning", "Findings", S, r, col, "Findings", "AG-FIND-MET-WEAKNESS", f'MET finding contains possible weakness or observation language: "{mm.group()}".',
                                             f"{GUIDE}: for MET, the narrative contains no advisory, observation, improvement, or weakness language that did not affect the determination.",
                                             text, None, False, obj, (body_start + mm.start(), body_start + mm.end())))
                        break
            elif status == "NOT MET":
                letter = re.search(r"\[([a-z])\]$", obj or "")
                refs = [obj] if obj else []
                if letter:
                    refs += [f"[{letter.group(1)}]", f"objective {letter.group(1)}", f"({letter.group(1)})"]
                if not any(x and x.lower() in text.lower() for x in refs):
                    out.append(self._new("Warning", "Findings", S, r, col, "Findings", "AG-FIND-NOTMET-OBJ", f"NOT MET finding does not cite the unsatisfied objective ({obj}).",
                                         f"{GUIDE}: for NOT MET, identify the exact unsatisfied objective.", text, None, False, obj))
                if not any(re.search(r"\b" + re.escape(term.strip()), low) for term in fcfg["not_met_deficiency_terms"]):
                    out.append(self._new("Warning", "Findings", S, r, col, "Findings", "AG-FIND-NOTMET-DEF", "NOT MET finding does not clearly state the factual evidence deficiency.",
                                         f"{GUIDE}: for NOT MET, identify the factual evidence deficiency (what was not provided, configured, or demonstrated).", text, None, False, obj))
            elif status == "NOT APPLICABLE":
                if len(body.split()) < 6 or not any(re.search(r"\b" + re.escape(term.strip()), low) for term in fcfg["na_reason_terms"]):
                    out.append(self._new("Warning", "Findings", S, r, col, "Findings", "AG-FIND-NA-REASON", "N/A finding does not explain why the objective does not apply.",
                                         f"{GUIDE}: for N/A, explain why the objective does not apply to the assessed environment.", text, None, False, obj))
            if len(body.split()) < 4 and status != "NOT APPLICABLE":
                out.append(self._new("Warning", "Findings", S, r, col, "Findings", "AG-FIND-THIN", "Findings contain a result label with little or no rationale.",
                                     f"{GUIDE}: follow the result with objective-specific factual rationale.", text, None, False, obj))
        if len(text) > fcfg.get("max_summary_chars", 1000):
            out.append(self._new("Info", "Findings", S, r, col, "Findings", "AG-FIND-LONG", f"Findings are {len(text):,} characters.",
                                 f"{GUIDE}: the Findings column is a short summary of the determination; detailed evidence belongs in Examine and Overall Comments. (Threshold set in rules.json.)",
                                 text, None, False, obj))
        out += self._text_issues(S, r, col, "Findings", text, obj)
        return out

    # ------------------------------------------------------------------
    def recheck_cell(self, sheet, row, col):
        """Re-run checks for one edited narrative cell; returns the new issue list for that cell."""
        self.issues = [i for i in self.issues if not (i.sheet == sheet and i.row == row and i.col == col and i.status != "Fixed")]
        new = []
        if sheet == S_OBJ:
            obj = next((o for rr, _q, o in self.res.obj_rows if rr == row), "")
            txt = as_text(self.res.get(sheet, row, col))
            if col == OC["findings"]:
                new = self.check_findings(row, obj, as_text(self.res.get(sheet, row, OC["score"])).strip())
            elif col in (OC["examine"], OC["test"], OC["overall"]):
                fname = OBJ_FIELD[col]
                if txt.strip():
                    new = self._limit(sheet, row, col, fname, txt, 4000, obj) + self._text_issues(sheet, row, col, fname, txt, obj)
            elif col == OC["esp"] and txt.strip():
                new = self._limit(sheet, row, col, "ESP Name", txt, 2000, obj)
        elif sheet == S_ASM:
            name = as_text(self.res.wb[S_ASM].cell(row, 1).value)
            txt = as_text(self.res.get(sheet, row, col))
            if name == "C3PAO Executive Summary":
                new = self._limit(sheet, row, col, name, txt, 2000) + self._text_issues(sheet, row, col, "Executive Summary", txt)
            elif "Professional Number" in name and txt and not re.fullmatch(r"\d+", txt.strip()):
                new = [self._new("Error", "Format", sheet, row, col, name, "EM-CPN", f'CPN "{txt}" must be the ID number only.', f'{INSTR} #10.', txt)]
        self.issues.extend(new)
        return new

    def accept(self, issue: Issue, text: str | None = None):
        """Apply an issue's proposed text (or user-edited text) to the working copy."""
        new = issue.proposed if text is None else text
        if new is None:
            raise ValueError("No proposed text for this issue.")
        if (issue.sheet, issue.col) not in EDITABLE:
            raise ValueError(PROTECTED_NOTE)
        cur = as_text(self.res.get(issue.sheet, issue.row, issue.col))
        if issue.field == "Findings" and issue.rule not in ("AG-FIND-OPENER", "AG-FIND-FORMAT"):
            a, b = FIND_RX.match(cur), FIND_RX.match(new)
            if a and (not b or norm_status(a.group(2)) != norm_status(b.group(2))):
                raise ValueError("This edit would change the Findings result (MET / NOT MET / N/A). The result is not modified during this pass.")
        self.res.set_text(issue.sheet, issue.row, issue.col, new, issue.rule)
        issue.status = "Fixed"
        return self.recheck_cell(issue.sheet, issue.row, issue.col)

    def ignore(self, issue: Issue):
        issue.status = "Ignored"
        self.ignored.add(issue.key())

    def apply_safe(self, issues=None, progress=None):
        """Apply all mechanical fixes (whitespace, doubled punctuation, term capitalization...). Loops per cell until stable."""
        cells = sorted({(i.sheet, i.row, i.col) for i in (issues or self.issues) if i.safe and i.status == "Open"})
        count = 0
        for n, (s, r, c) in enumerate(cells):
            for _ in range(40):
                cand = [i for i in self.issues if (i.sheet, i.row, i.col) == (s, r, c) and i.safe and i.status == "Open" and i.proposed is not None]
                if not cand:
                    break
                try:
                    self.accept(cand[0])
                    count += 1
                except ValueError:
                    cand[0].safe = False
            if progress:
                progress(n + 1, len(cells))
        return count

    def summary(self):
        open_ = [i for i in self.issues if i.status == "Open"]
        by = lambda k: {x: sum(1 for i in open_ if getattr(i, k) == x) for x in sorted({getattr(i, k) for i in open_})}  # noqa: E731
        return {"total": len(self.issues), "open": len(open_), "fixed": sum(i.status == "Fixed" for i in self.issues),
                "ignored": sum(i.status == "Ignored" for i in self.issues), "by_severity": by("severity"), "by_category": by("category"), "by_rule": by("rule")}
