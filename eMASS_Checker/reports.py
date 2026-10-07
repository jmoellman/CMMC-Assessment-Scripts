"""Issue reports: printable HTML, Excel, and CSV."""
from __future__ import annotations

import csv
import datetime as dt
import hashlib
import html
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SEV_ORDER = {"Error": 0, "Warning": 1, "Info": 2}
COLUMNS = ["#", "Status", "Severity", "Sheet", "Cell", "Row", "Column", "Field", "Objective", "Rule", "Category", "Issue", "Why it is an issue", "Flagged text", "Current cell text", "Proposed fix"]


def file_sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _flagged(i):
    a, b = i.span
    if (a, b) == (0, 0) or not i.current:
        return ""
    return i.current[a:b] if b > a else "(end of text)"


def issue_rows(issues, include_closed=True):
    rows = []
    for n, i in enumerate(sorted(issues, key=lambda x: (x.sheet, x.row, x.col, SEV_ORDER[x.severity])), start=1):
        if not include_closed and i.status != "Open":
            continue
        rows.append([n, i.status, i.severity, i.sheet, i.cell, i.row or "", i.col_letter, i.field, i.objective, i.rule, i.category, i.message, i.why,
                     _flagged(i), i.current, i.proposed if i.proposed is not None else ""])
    return rows


def write_csv(path, issues, include_closed=True):
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(COLUMNS)
        w.writerows(issue_rows(issues, include_closed))


def write_xlsx(path, engine, include_closed=True):
    wb = Workbook()
    ws = wb.active
    ws.title = "Issues"
    ws.append(["***** CUI (When Filled In) *****"])
    ws["A1"].font = Font(bold=True, color="C00000")
    meta = report_meta(engine)
    for k, v in meta.items():
        ws.append([k, v])
    ws.append([])
    hdr_row = ws.max_row + 1
    ws.append(COLUMNS)
    for c in ws[hdr_row]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="292C35")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    fills = {"Error": "F8D7DA", "Warning": "FFF3CD", "Info": "D9EAF7"}
    for r in issue_rows(engine.issues, include_closed):
        ws.append(r)
        row = ws.max_row
        ws.cell(row, 3).fill = PatternFill("solid", fgColor=fills.get(r[2], "FFFFFF"))
        for c in ws[row]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    widths = [5, 9, 9, 20, 7, 6, 7, 18, 18, 22, 16, 50, 55, 25, 60, 60]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(hdr_row + 1, 1)
    ws.auto_filter.ref = f"A{hdr_row}:{get_column_letter(len(COLUMNS))}{ws.max_row}"
    # change log
    cl = wb.create_sheet("Change Log")
    cl.append(["Time", "Location", "Rule / reason", "Before", "After"])
    for c in cl[1]:
        c.font = Font(bold=True)
    for e in engine.res.change_log:
        cl.append([e["time"], e["location"], e["rule"], e["before"], e["after"]])
    for i, w in enumerate([18, 28, 24, 70, 70], start=1):
        cl.column_dimensions[get_column_letter(i)].width = w
    for row in cl.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    # summary
    sm = wb.create_sheet("Summary")
    s = engine.summary()
    sm.append(["Open issues by severity"])
    for k, v in sorted(s["by_severity"].items(), key=lambda kv: SEV_ORDER[kv[0]]):
        sm.append([k, v])
    sm.append([])
    sm.append(["Open issues by rule"])
    for k, v in sorted(s["by_rule"].items(), key=lambda kv: -kv[1]):
        sm.append([k, v])
    sm.column_dimensions["A"].width = 30
    wb.save(path)


def report_meta(engine):
    r = engine.res
    s = engine.summary()
    return {
        "Results file": os.path.basename(r.path),
        "File SHA-256 (as loaded)": file_sha256(r.path),
        "Template version": r.template_version,
        "Scope": engine.scope.describe(),
        "Report generated": dt.datetime.now().strftime("%d-%b-%Y %H:%M"),
        "Issues": f"{s['open']} open, {s['fixed']} fixed, {s['ignored']} ignored (total {s['total']})",
        "Cells changed in working copy": str(len(r.changes)),
        "Grammar engine": "Built-in rules + offline dictionary" + (" + LanguageTool (local)" if engine.opts.get("languagetool") else ""),
    }


def _hl(i):
    """Current text with the flagged span highlighted (HTML-escaped)."""
    t = i.current or ""
    a, b = i.span
    if not t:
        return ""
    if (a, b) == (0, 0) or a > len(t):
        body = html.escape(t)
    else:
        mark = html.escape(t[a:b]) if b > a else "&#8203;&#9650;"
        body = html.escape(t[:a]) + f"<mark>{mark}</mark>" + html.escape(t[b:])
    return body.replace("\n", "<br>")


def write_html(path, engine, include_closed=False, auto_print=False):
    meta = report_meta(engine)
    s = engine.summary()
    issues = sorted([i for i in engine.issues if include_closed or i.status == "Open"], key=lambda x: (x.sheet, x.row, x.col, SEV_ORDER[x.severity]))
    sev_cards = "".join(f'<div class="card {k.lower()}"><b>{s["by_severity"].get(k, 0)}</b><span>{k}</span></div>' for k in ("Error", "Warning", "Info"))
    rule_rows = "".join(f'<span class="chip">{html.escape(k)} <b>{v}</b></span>' for k, v in sorted(s["by_rule"].items(), key=lambda kv: -kv[1]))
    meta_rows = "".join(f"<tr><th>{html.escape(k)}</th><td>{html.escape(v)}</td></tr>" for k, v in meta.items())
    body = []
    for n, i in enumerate(issues, start=1):
        prop = f'<div class="prop"><b>Proposed fix:</b> {html.escape(i.proposed)}</div>' if i.proposed is not None and i.status == "Open" else ""
        cur = f'<div class="cur">{_hl(i)}</div>' if i.current else ""
        body.append(
            f'<tr class="{i.severity.lower()}"><td>{n}</td><td><span class="sev {i.severity.lower()}">{i.severity}</span><br><small>{i.status}</small></td>'
            f'<td><b>{html.escape(i.sheet)}</b><br>Cell {html.escape(i.cell)}<br>Row {i.row or "-"} / Col {html.escape(i.col_letter or "-")}</td>'
            f'<td>{html.escape(i.field)}<br><small>{html.escape(i.objective)}</small></td>'
            f'<td><b>{html.escape(i.message)}</b><div class="why"><b>Why:</b> {html.escape(i.why)}</div>{cur}{prop}<div class="rule">{html.escape(i.rule)} &middot; {html.escape(i.category)}</div></td></tr>')
    changes = "".join(f"<tr><td>{html.escape(e['location'])}</td><td>{html.escape(e['rule'])}</td><td>{html.escape(e['before'])}</td><td>{html.escape(e['after'])}</td></tr>"
                      for e in engine.res.change_log)
    change_section = (f"<h2>Change log ({len(engine.res.change_log)})</h2><table class='log'><tr><th>Cell</th><th>Reason</th><th>Before</th><th>After</th></tr>{changes}</table>"
                      if engine.res.change_log else "")
    doc = f"""<!doctype html><html><head><meta charset="utf-8"><title>eMASS Results Check Report</title>
<style>
body{{font-family:Calibri,Segoe UI,Arial,sans-serif;color:#222;margin:24px;font-size:12.5px}}
.banner{{text-align:center;font-weight:bold;color:#b00000;letter-spacing:1px;margin-bottom:8px}}
h1{{font-size:20px;color:#292C35;margin:6px 0}} h2{{font-size:15px;color:#292C35;border-bottom:2px solid #F7941D;padding-bottom:3px;margin-top:22px}}
table{{border-collapse:collapse;width:100%}} th,td{{border:1px solid #ccc;padding:5px 7px;vertical-align:top;text-align:left}}
.meta th{{width:220px;background:#f3f3f5}} .cards{{display:flex;gap:10px;margin:10px 0}}
.card{{border-radius:6px;padding:8px 16px;min-width:90px;text-align:center}} .card b{{display:block;font-size:22px}}
.card.error{{background:#f8d7da}} .card.warning{{background:#fff3cd}} .card.info{{background:#d9eaf7}}
.sev{{padding:1px 6px;border-radius:3px;font-weight:bold;font-size:11px}} .sev.error{{background:#dc3545;color:#fff}} .sev.warning{{background:#ffc107}} .sev.info{{background:#5b9bd5;color:#fff}}
.why{{margin-top:4px;color:#444}} .cur{{margin-top:5px;background:#fafafa;border-left:3px solid #bbb;padding:4px 6px;white-space:pre-wrap}}
.prop{{margin-top:5px;background:#eef8ee;border-left:3px solid #4a9a4a;padding:4px 6px;white-space:pre-wrap}} mark{{background:#ffe08a}}
.chips{{line-height:2}} .chip{{display:inline-block;border:1px solid #ccc;border-radius:12px;padding:0 8px;margin:0 4px 0 0;font-size:11px;background:#f6f6f8}}
.rule{{margin-top:4px;color:#888;font-size:11px}} .issues thead th{{background:#292C35;color:#fff}}
tr{{page-break-inside:avoid}} @media print{{body{{margin:8mm}} .noprint{{display:none}} thead{{display:table-header-group}}}}
</style></head><body>
<div class="banner">***** CUI (When Filled In) *****</div>
<h1>CMMC eMASS Assessment Results: Check Report</h1>
<p class="noprint"><button onclick="window.print()">Print</button></p>
<table class="meta">{meta_rows}</table>
<h2>Summary of open issues</h2><div class="cards">{sev_cards}</div>
<div class="chips"><b>Open issues by rule:</b> {rule_rows}</div>
<h2>{'All' if include_closed else 'Open'} issues ({len(issues)})</h2>
<table class="issues"><thead><tr><th>#</th><th>Severity</th><th>Location</th><th>Field</th><th>Issue, reason, and cell text</th></tr></thead><tbody>{''.join(body)}</tbody></table>
{change_section}
<div class="banner" style="margin-top:16px">***** CUI (When Filled In) *****</div>
{'<script>window.onload=function(){window.print();}</script>' if auto_print else ''}
</body></html>"""
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(doc)
